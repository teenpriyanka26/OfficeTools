from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import ExtractedTable


_INVALID_SHEET_CHARS = re.compile(r"[\\/?*:[\]]")


def _safe_sheet_name(title: str, index: int, used_names: set[str]) -> str:
    base = re.sub(r"\s+", " ", _INVALID_SHEET_CHARS.sub(" ", title or "")).strip()
    base = (base or f"Table {index}")[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used_names:
        marker = f" ({suffix})"
        suffix += 1
        candidate = f"{base[: 31 - len(marker)]}{marker}"
    used_names.add(candidate.casefold())
    return candidate


def export_tables(tables: list[ExtractedTable], output_path: str | Path) -> Path:
    """Create an Excel workbook using the Python-only openpyxl runtime."""
    if not tables:
        raise ValueError("No tables were extracted; no workbook was created")

    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()

    for index, table in enumerate(tables, start=1):
        sheet = workbook.create_sheet(_safe_sheet_name(table.title, index, used_names))
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if not table.rows:
            continue

        width = max(len(row) for row in table.rows)
        rows = [(row + [""] * width)[:width] for row in table.rows]
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = sheet.cell(row=row_index, column=column_index, value=str(value))
                # Extracted values are source text, including strings beginning with '='.
                cell.data_type = "s"
                cell.font = Font(name="Aptos", size=10, color="1F2937", bold=row_index == 1)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.fill = PatternFill("solid", fgColor="DCE6F1")

        for column_index in range(1, width + 1):
            longest = max(len(str(row[column_index - 1])) for row in rows)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                36, max(10, longest + 2)
            )

        end_column = get_column_letter(width)
        excel_table = Table(
            displayName=f"ExtractedTable{index}",
            ref=f"A1:{end_column}{len(rows)}",
        )
        excel_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        sheet.add_table(excel_table)

    workbook.save(output)
    return output
