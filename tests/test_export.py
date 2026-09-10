from openpyxl import load_workbook

from office_toolkit.export import export_tables
from office_toolkit.models import ExtractedTable


def test_exports_tables_and_preserves_source_text(tmp_path):
    output = tmp_path / "result.xlsx"
    tables = [
        ExtractedTable(
            title="Assets/Current",
            rows=[["ID", "Value"], ["1", "=not-a-formula"]],
            source_pages=[1],
        ),
        ExtractedTable(
            title="assets current",
            rows=[["Role", "Name"], ["Manager", "A. Example"]],
            source_pages=[2],
        ),
    ]

    result = export_tables(tables, output)

    assert result == output.resolve()
    workbook = load_workbook(result, data_only=False)
    assert workbook.sheetnames == ["Assets Current", "assets current (2)"]
    assert workbook["Assets Current"]["B2"].value == "=not-a-formula"
    assert workbook["Assets Current"]["B2"].data_type == "s"
    assert workbook["Assets Current"].freeze_panes == "A2"
    assert not workbook["Assets Current"].sheet_view.showGridLines
    assert len(workbook["Assets Current"].tables) == 1


def test_rejects_empty_table_collection(tmp_path):
    try:
        export_tables([], tmp_path / "empty.xlsx")
    except ValueError as error:
        assert str(error) == "No tables were extracted; no workbook was created"
    else:
        raise AssertionError("Expected export_tables() to reject an empty collection")
